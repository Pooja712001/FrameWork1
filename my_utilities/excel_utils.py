import openpyxl

def get_data_from_excel(file_name, sheet_name):
    workbook = openpyxl.load_workbook(file_name)
    sheet = workbook[sheet_name]
    data = []
    for row in range(2, sheet.max_row + 1):
        email = sheet.cell(row, 1).value
        password = sheet.cell(row, 2).value
        expected = sheet.cell(row, 3).value
        data.append((email, password, expected))
    return data
