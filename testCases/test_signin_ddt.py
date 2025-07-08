import pytest
from selenium import webdriver
from my_utilities.read_config import ReadConfig
from Pages.Signin_page import SigninPage
from my_utilities.logger import log
import openpyxl
import os

# Read data from Excel
def get_login_data():
    file_path = os.path.join("TestData", "login_data.xlsx")
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook["Sheet1"]
    data = []

    for row in range(2, sheet.max_row + 1):  # Skip header
        email = sheet.cell(row, 1).value
        password = sheet.cell(row, 2).value
        expected = sheet.cell(row, 3).value.upper()  # Ensure PASS/FAIL is uppercase
        data.append((email, password, expected))
    return data

@pytest.mark.parametrize("email,password,expected", get_login_data())
def test_sign_in_ddt(setup, email, password, expected):
    driver = setup
    driver.get(ReadConfig.get_base_url())
    log.info(f"Testing login with Email: {email} | Expected: {expected}")

    login = SigninPage(driver)
    login.enter_email(email)
    login.enter_password(password)
    login.click_sign_in()

    try:
        account_name = login.get_account_name()
        if expected == "PASS":
            log.info("Login successful as expected.")
            assert True
        else:
            log.error("Login succeeded but expected failure.")
            assert False
    except:
        if expected == "FAIL":
            log.info("Login failed as expected.")
            assert True
        else:
            log.error("Login failed but expected success.")
            assert False
